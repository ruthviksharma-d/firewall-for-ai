"""
File Scanner - extracts plain text from uploaded files so it can be run
through EXACTLY the same detector pipeline as a typed prompt (no separate
detection logic is duplicated here; this module's only job is text
extraction).

Supported formats: PDF, DOCX, CSV, XLSX, TXT, PNG, JPEG/JPG.
Image OCR requires the system `tesseract` binary (apt-get install
tesseract-ocr) - if it's missing, extraction fails gracefully with a clear
reason instead of crashing the scan.
"""
import base64
import csv
import io
import logging

logger = logging.getLogger("promptshield.ai.file_scanner")

SUPPORTED_EXTENSIONS = {"pdf", "docx", "csv", "xlsx", "txt", "png", "jpg", "jpeg"}


class FileExtractionResult:
    def __init__(self, filename: str, text: str, success: bool, reason: str = ""):
        self.filename = filename
        self.text = text
        self.success = success
        self.reason = reason


def extract_text_from_file(filename: str, content_base64: str) -> FileExtractionResult:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension not in SUPPORTED_EXTENSIONS:
        return FileExtractionResult(filename, "", False, f"Unsupported file type: .{extension}")

    try:
        raw = base64.b64decode(content_base64, validate=False)
    except Exception as exc:
        return FileExtractionResult(filename, "", False, f"Invalid base64 content: {exc}")

    try:
        if extension == "pdf":
            text = _extract_pdf(raw)
        elif extension == "docx":
            text = _extract_docx(raw)
        elif extension == "csv":
            text = _extract_csv(raw)
        elif extension == "xlsx":
            text = _extract_xlsx(raw)
        elif extension == "txt":
            text = raw.decode("utf-8", errors="ignore")
        elif extension in {"png", "jpg", "jpeg"}:
            text = _extract_image_ocr(raw)
        else:
            return FileExtractionResult(filename, "", False, f"Unsupported file type: .{extension}")
    except Exception as exc:
        logger.warning("Failed to extract text from %s: %s", filename, exc)
        return FileExtractionResult(filename, "", False, f"Extraction failed: {exc}")

    return FileExtractionResult(filename, text.strip(), True)


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(raw: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_csv(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(", ".join(row) for row in reader)


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                lines.append(", ".join(values))
    return "\n".join(lines)


def _extract_image_ocr(raw: bytes) -> str:
    import pytesseract
    from PIL import Image

    image = Image.open(io.BytesIO(raw))
    return pytesseract.image_to_string(image)
